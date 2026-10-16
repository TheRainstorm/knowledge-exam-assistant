"""Local documents/video -> verbatim chunks, page/slide images and source references."""
from pathlib import Path
import json
import shutil
import subprocess
import tempfile

from PIL import Image
import numpy as np

from .storage import atomic_json, digest_file, fingerprint, writer_lock

SUPPORTED = {".md", ".txt", ".pdf", ".docx", ".pptx", ".png", ".jpg", ".jpeg", ".webp", ".mp4", ".mkv", ".mov", ".webm"}
VIDEO = {".mp4", ".mkv", ".mov", ".webm"}


def chunks(text, size=1400, overlap=160):
    # Contiguous windows retain definitions/numbers; never paraphrase evidence.
    text = text.replace("\r\n", "\n").strip()
    for start in range(0, len(text), size - overlap):
        block = text[start:start + size]
        if block.strip():
            yield block
        if start + size >= len(text):
            break


def ingest(cfg, source_dir):
    root, source_dir = Path(cfg["data_dir"]), Path(source_dir).resolve()
    if not source_dir.is_dir():
        raise ValueError("Input must be an existing directory")
    if root.is_relative_to(source_dir):
        raise ValueError("Keep data_dir outside the input directory to avoid indexing generated files")
    settings = cfg["ingest"]
    ocr_engine = None
    if settings.get("ocr"):
        from rapidocr_onnxruntime import RapidOCR
        ocr_engine = RapidOCR()
    asr_engine = None
    if settings.get("transcribe"):
        from faster_whisper import WhisperModel
        asr_engine = WhisperModel(settings.get("asr_model", "small"),
            device=settings.get("asr_device", "cpu"), compute_type=settings.get("asr_compute_type", "int8"))

    with writer_lock(root):
        manifest_path = root / "corpus.json"
        # Each command publishes a complete inventory of THIS input directory.
        # Prior manifests remain active if any source fails before final publication.
        records, sources = [], []
        for path in sorted(source_dir.rglob("*")):
            if not path.is_file() or path.is_symlink() or path.suffix.lower() not in SUPPORTED:
                continue
            name = str(path.relative_to(source_dir)).replace("\\", "/")
            source_sha = digest_file(path)
            source_id = "s_" + fingerprint({"name": name, "sha256": source_sha})[:20]
            signature = fingerprint({"schema": 1, "source_sha": source_sha, "settings": settings, "name": name})
            cache = root / "cache/ingest" / (signature + ".json")
            asset_dir = root / "outputs/deliverables/assets" / source_id
            if cache.is_file():
                rows = json.loads(cache.read_text(encoding="utf-8"))
                valid = all(not r.get("image") or ((root / "outputs/deliverables" / r["image"]).is_file()
                            and digest_file(root / "outputs/deliverables" / r["image"]) == r["image_sha256"]) for r in rows)
            else:
                valid = False
            if valid:
                print("reuse", name, flush=True)
            else:
                asset_dir.mkdir(parents=True, exist_ok=True)
                rows = extract(path, source_id, name, asset_dir, settings, ocr_engine, asr_engine)
                atomic_json(cache, rows)
                print("ingest", name, len(rows), flush=True)
            records.extend(rows)
            sources.append({"id": source_id, "name": name, "sha256": source_sha,
                            "path": str(path), "kind": "video" if path.suffix.lower() in VIDEO else "document"})
        if not sources:
            raise ValueError("No supported input files found")
        for n, record in enumerate(records, 1):
            record["record_id"] = f"r{n}"
            if record["kind"] != "ppt":
                record["handbook_url"] = f"/handbook#r{n}"
        corpus = {"schema_version": 1, "sources": sources, "records": records}
        corpus["fingerprint"] = fingerprint(corpus)
        atomic_json(manifest_path, corpus)
        # A readable, deterministic evidence text; no unsupported AI summaries.
        text = ["# 知识原文汇编", "", "自动提取原文，OCR/ASR 可能有误；请核对来源。", ""]
        for r in records:
            text.extend([f"## {r['record_id']} · {r['course']} · {r.get('section', '')}", "", r["text"], ""])
        (root / "knowledge.md").write_text("\n".join(text), encoding="utf-8")
        return {"sources": len(sources), "records": len(records), "fingerprint": corpus["fingerprint"]}


def extract(path, sid, name, asset_dir, settings, ocr_engine, asr_engine):
    rows = []

    def add(text, kind="handbook", **extra):
        for block in chunks(text):
            rows.append({"source_id": sid, "course_id": sid, "course": name, "kind": kind,
                         "text": block, **extra})

    def image_row(image, filename, text="", **extra):
        target = asset_dir / filename
        image.convert("RGB").save(target, format="PNG")
        if ocr_engine is not None:
            result, _ = ocr_engine(str(target))
            recognized = "\n".join(str(line[1]) for line in result or [])
            if recognized and recognized not in text:
                text = (text + "\n" + recognized).strip()
        rows.append({"source_id": sid, "course_id": sid, "course": name, "kind": "ppt",
                     "image": f"assets/{sid}/{filename}", "image_sha256": digest_file(target),
                     "text": text, "ocr_text": text, **extra})

    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        add(path.read_text(encoding="utf-8-sig"))
    elif suffix == ".docx":
        from docx import Document
        document = Document(path)
        add("\n\n".join(p.text for p in document.paragraphs))
        for table in document.tables:
            add("\n".join(" | ".join(c.text for c in row.cells) for row in table.rows))
    elif suffix == ".pptx":
        if not shutil.which("libreoffice"):
            raise RuntimeError("PPTX import requires LibreOffice, or export PPTX to PDF first")
        with tempfile.TemporaryDirectory() as temp:
            profile = Path(temp) / "lo-profile"
            subprocess.run(["libreoffice", "-env:UserInstallation=" + profile.as_uri(), "--headless", "--convert-to", "pdf", "--outdir", temp, str(path)], check=True, timeout=300, capture_output=True)
            return extract(Path(temp) / (path.stem + ".pdf"), sid, name, asset_dir, settings, ocr_engine, asr_engine)
    elif suffix == ".pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            for n, page in enumerate(doc, 1):
                text = page.get_text()
                pix = page.get_pixmap(dpi=int(settings.get("pdf_dpi", 120)), alpha=False)
                image_row(Image.frombytes("RGB", (pix.width, pix.height), pix.samples), f"page_{n:05}.png", text, page=n, section=f"第 {n} 页")
                if text.strip():
                    add(text, page=n, section=f"第 {n} 页")
    elif suffix in VIDEO:
        import cv2
        cap = cv2.VideoCapture(str(path))
        fps, count = cap.get(cv2.CAP_PROP_FPS), cap.get(cv2.CAP_PROP_FRAME_COUNT)
        if not cap.isOpened() or fps <= 0 or count <= 0:
            cap.release()
            raise ValueError("Cannot decode video: " + name)
        duration, interval, anchor = count / fps, float(settings.get("sample_seconds", 5)), None
        roi = settings.get("compare_roi", [0, 0, 1, 1])
        if len(roi) != 4 or not (0 <= roi[0] < roi[2] <= 1 and 0 <= roi[1] < roi[3] <= 1):
            cap.release()
            raise ValueError("compare_roi must be normalized [left, top, right, bottom]")
        try:
            for n, stamp in enumerate(np.arange(0, duration, interval)):
                cap.set(cv2.CAP_PROP_POS_MSEC, float(stamp * 1000))
                ok, frame = cap.read()
                if not ok:
                    raise ValueError(f"Video decode failed at {stamp:.1f}s: {name}")
                gray = cv2.cvtColor(cv2.resize(frame, (320, 180)), cv2.COLOR_BGR2GRAY)
                crop = gray[int(roi[1]*180):max(int(roi[3]*180),1), int(roi[0]*320):max(int(roi[2]*320),1)]
                if anchor is not None and float(np.abs(crop.astype(float)-anchor).mean()) < settings.get("frame_change_threshold", 6):
                    continue
                anchor = crop.astype(float)
                image_row(Image.fromarray(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)), f"slide_{n:06}.png", start_s=round(float(stamp), 3),
                          section=f"{stamp:.1f} 秒", watch_url=f"/watch/{sid}?t={stamp:.1f}")
        finally:
            cap.release()
        if asr_engine is not None:
            with tempfile.TemporaryDirectory() as temp:
                wav = Path(temp) / "audio.wav"
                subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-i", str(path), "-vn", "-ac", "1", "-ar", "16000", str(wav)], check=True)
                segments, _ = asr_engine.transcribe(str(wav), language=settings.get("asr_language", "zh"), vad_filter=True)
                for segment in segments:
                    add(segment.text, "transcript", start_s=segment.start, end_s=segment.end,
                        watch_url=f"/watch/{sid}?t={segment.start:.1f}", section=f"{segment.start:.1f}–{segment.end:.1f} 秒")
    else:
        with Image.open(path) as image:
            image_row(image, "image.png")
    return rows
